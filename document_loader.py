import os
import zlib
import zipfile
import urllib.parse
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from typing import Generator, List
from pypdf import PdfReader
from pptx import Presentation
import olefile


@dataclass
class DocumentChunk:
    file_name: str
    file_path: str
    page_number: int
    text: str
    chunk_type: str


class DocumentScanner:
    SUPPORTED_EXTENSIONS = {'.pdf', '.pptx', '.hwp', '.hwpx'}

    def __init__(self, root_dir: str):
        self.root_dir = root_dir

    def scan_files(self) -> Generator[str, None, None]:
        for root, _, files in os.walk(self.root_dir):
            for file in files:
                if file.startswith('~$'):
                    continue
                ext = os.path.splitext(file)[1].lower()
                if ext in self.SUPPORTED_EXTENSIONS:
                    yield os.path.join(root, file)


class DocumentLoader:
    def _get_clean_filename(self, file_path: str) -> str:
        raw_name = os.path.basename(file_path)
        return urllib.parse.unquote(raw_name)

    def load_pdf(self, file_path: str) -> List[DocumentChunk]:
        chunks: List[DocumentChunk] = []
        file_name = self._get_clean_filename(file_path)

        try:
            reader = PdfReader(file_path)
            for page_idx, page in enumerate(reader.pages, start=1):
                try:
                    text = page.extract_text() or ""
                    clean_text = text.strip()
                    if clean_text:
                        chunks.append(DocumentChunk(
                            file_name=file_name,
                            file_path=file_path,
                            page_number=page_idx,
                            text=clean_text,
                            chunk_type="pdf"
                        ))
                except Exception:
                    continue
        except Exception:
            return []

        return chunks

    def load_pptx(self, file_path: str) -> List[DocumentChunk]:
        chunks: List[DocumentChunk] = []
        file_name = self._get_clean_filename(file_path)

        try:
            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides, start=1):
                slide_texts: List[str] = []

                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            line = paragraph.text.strip()
                            if line:
                                slide_texts.append(line)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            for cell in row.cells:
                                cell_text = cell.text.strip()
                                if cell_text:
                                    slide_texts.append(cell_text)

                combined_text = "\n".join(slide_texts).strip()
                if combined_text:
                    chunks.append(DocumentChunk(
                        file_name=file_name,
                        file_path=file_path,
                        page_number=slide_idx,
                        text=combined_text,
                        chunk_type="pptx"
                    ))
        except Exception:
            return []

        return chunks

    def load_hwpx(self, file_path: str) -> List[DocumentChunk]:
        chunks: List[DocumentChunk] = []
        file_name = self._get_clean_filename(file_path)

        try:
            with zipfile.ZipFile(file_path, 'r') as z:
                section_files = [f for f in z.namelist() if f.startswith('Contents/section') and f.endswith('.xml')]
                section_files.sort()

                for page_idx, s_file in enumerate(section_files, start=1):
                    xml_data = z.read(s_file)
                    root = ET.fromstring(xml_data)
                    texts = [elem.text.strip() for elem in root.iter() if elem.text and elem.text.strip()]
                    combined = "\n".join(texts)

                    if combined:
                        chunks.append(DocumentChunk(
                            file_name=file_name,
                            file_path=file_path,
                            page_number=page_idx,
                            text=combined,
                            chunk_type="hwpx"
                        ))
        except Exception:
            return []

        return chunks

    def load_hwp(self, file_path: str) -> List[DocumentChunk]:
        chunks: List[DocumentChunk] = []
        file_name = self._get_clean_filename(file_path)

        try:
            if not olefile.isOleFile(file_path):
                return []

            with olefile.OleFileIO(file_path) as ole:
                dirs = ole.listdir()
                section_dirs = [d for d in dirs if len(d) > 1 and d[0] == 'BodyText' and d[1].startswith('Section')]
                section_dirs.sort(key=lambda x: int(x[1].replace('Section', '')) if x[1].replace('Section', '').isdigit() else 0)

                for page_idx, sec in enumerate(section_dirs, start=1):
                    stream = ole.openstream(sec).read()
                    try:
                        decompressed = zlib.decompress(stream, -15)
                    except Exception:
                        decompressed = stream

                    text = decompressed.decode('utf-16le', errors='ignore')
                    clean_lines = [line.strip() for line in text.splitlines() if line.strip()]
                    combined = "\n".join(clean_lines)

                    if combined:
                        chunks.append(DocumentChunk(
                            file_name=file_name,
                            file_path=file_path,
                            page_number=page_idx,
                            text=combined,
                            chunk_type="hwp"
                        ))
        except Exception:
            return []

        return chunks

    def process_file(self, file_path: str) -> List[DocumentChunk]:
        if not os.path.exists(file_path) or os.path.getsize(file_path) == 0:
            return []

        ext = os.path.splitext(file_path)[1].lower()
        if ext == '.pdf':
            return self.load_pdf(file_path)
        elif ext == '.pptx':
            return self.load_pptx(file_path)
        elif ext == '.hwpx':
            return self.load_hwpx(file_path)
        elif ext == '.hwp':
            return self.load_hwp(file_path)
        return []


if __name__ == "__main__":
    test_folder = "./test_docs"
    loader = DocumentLoader()
    scanner = DocumentScanner(test_folder)

    if not os.path.exists(test_folder):
        os.makedirs(test_folder, exist_ok=True)
        print(f"폴더 생성: {test_folder}")
    else:
        for file_path in scanner.scan_files():
            results = loader.process_file(file_path)
            clean_name = urllib.parse.unquote(os.path.basename(file_path))
            print(f"\n{clean_name} ({len(results)} chunks)")
            for chunk in results[:2]:
                preview = chunk.text[:70].replace('\n', ' ')
                print(f"  [{chunk.chunk_type.upper()}] p.{chunk.page_number}: {preview}...")
